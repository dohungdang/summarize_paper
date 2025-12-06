import gradio as gr
import torch
from transformers import AutoTokenizer, AutoModelForSeq2SeqLM
import docx
import PyPDF2
import re
import requests
from newspaper import Article
from pathlib import Path

DEFAULT_MODEL_DIR = "outputs/bartpho-finetuned"
DEFAULT_TEST_DOCX = "/mnt/data/baocaohocmay.docx"


def read_pdf(file_obj):
    reader = PyPDF2.PdfReader(file_obj)
    text = []
    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            text.append(page_text)
    return "\n".join(text)


def read_docx(path_or_file):
    if hasattr(path_or_file, "read"):
        doc = docx.Document(path_or_file)
    else:
        doc = docx.Document(str(path_or_file))
    return "\n".join([p.text for p in doc.paragraphs if p.text.strip()])


def clean_text(text: str) -> str:
    if not text:
        return ""
    text = text.strip()
    text = re.sub(r"\t+", " ", text)
    text = re.sub(r" *\n+ *", "\n", text)
    text = re.sub(r" {2,}", " ", text)
    return text


# Cache model
tokenizer_cache = None
model_cache = None
device_cache = None


def load_model_and_tokenizer(model_dir: str):
    global tokenizer_cache, model_cache, device_cache
    if tokenizer_cache is not None:
        return tokenizer_cache, model_cache, device_cache

    tokenizer_cache = AutoTokenizer.from_pretrained(model_dir)
    model_cache = AutoModelForSeq2SeqLM.from_pretrained(model_dir)
    device_cache = "cuda" if torch.cuda.is_available() else "cpu"
    model_cache = model_cache.to(device_cache)
    return tokenizer_cache, model_cache, device_cache


def chunk_text_by_tokens(text: str, tokenizer, max_tokens: int = 800, overlap: int = 64):
    ids = tokenizer.encode(text)
    if len(ids) <= max_tokens:
        return [text]
    chunks = []
    start = 0
    while start < len(ids):
        end = start + max_tokens
        chunk_ids = ids[start:end]
        chunk_text = tokenizer.decode(chunk_ids, skip_special_tokens=True, clean_up_tokenization_spaces=True)
        chunk_text = clean_text(chunk_text)
        if chunk_text:
            chunks.append(chunk_text)
        start = max(0, end - overlap)
        if end >= len(ids):
            break
    return chunks


def summarize_chunk(chunk: str, tokenizer, model, device, max_summary_tokens=128, num_beams=4):
    inputs = tokenizer(chunk, return_tensors="pt", truncation=True, max_length=1024)
    input_ids = inputs["input_ids"].to(device)
    attention_mask = inputs.get("attention_mask", None)

    gen_kwargs = dict(max_length=max_summary_tokens, num_beams=num_beams, early_stopping=True)

    if attention_mask is not None:
        gen = model.generate(input_ids=input_ids, attention_mask=attention_mask.to(device), **gen_kwargs)
    else:
        gen = model.generate(input_ids=input_ids, **gen_kwargs)

    summary = tokenizer.decode(gen[0], skip_special_tokens=True)
    return clean_text(summary)


def extract_text_from_url(url: str, timeout: int = 10):
    try:
        article = Article(url)
        article.download()
        article.parse()
        txt = clean_text(article.text)
        if txt and len(txt) > 50:
            return txt
    except Exception:
        pass

    try:
        r = requests.get(url, timeout=timeout)
        html = r.text
        text = re.sub(r"<script.*?>.*?</script>", "", html, flags=re.S)
        text = re.sub(r"<style.*?>.*?</style>", "", text, flags=re.S)
        text = re.sub(r"<[^>]+>", " ", text)
        return clean_text(text)
    except Exception as e:
        raise RuntimeError(f"Không thể tải nội dung từ URL: {e}")


# ============================================================
# ======================  GRADIO APP  ========================
# ============================================================

def process_input(model_dir, max_chunk_tokens, chunk_overlap, max_summary_tokens,
                  num_beams, file, url, use_sample):

    # Load model
    tokenizer, model, device = load_model_and_tokenizer(model_dir)

    text = ""
    source_label = ""

    # File input
    if file is not None:
        source_label = file.name
        if file.name.endswith(".pdf"):
            text = read_pdf(file)
        else:
            text = read_docx(file)

    # URL input
    elif url:
        source_label = url
        text = extract_text_from_url(url)

    # Sample file
    elif use_sample:
        path = Path(DEFAULT_TEST_DOCX)
        if path.exists():
            source_label = str(path)
            text = read_docx(path)
        else:
            return f"File mẫu không tồn tại: {path}", "", ""

    else:
        return "Vui lòng cung cấp PDF / DOCX hoặc URL", "", ""

    text = clean_text(text)

    # Chunking
    chunks = chunk_text_by_tokens(text, tokenizer, max_tokens=max_chunk_tokens, overlap=chunk_overlap)

    summaries = []
    for ch in chunks:
        s = summarize_chunk(ch, tokenizer, model, device,
                            max_summary_tokens=max_summary_tokens,
                            num_beams=num_beams)
        summaries.append(s)

    final_summary = "\n\n".join(summaries)

    # Short summary
    short_summary = summarize_chunk(final_summary, tokenizer, model, device,
                                    max_summary_tokens=200, num_beams=4)

    return text[:3000], final_summary, short_summary


with gr.Blocks(title="BartPho Summarizer") as demo:
    gr.Markdown("# **BartPho — Summarizer (Gradio UI)**")

    with gr.Row():
        with gr.Column():
            model_dir = gr.Textbox(value=DEFAULT_MODEL_DIR, label="Model directory")
            file = gr.File(label="Upload PDF / DOCX")
            url = gr.Textbox(label="Hoặc nhập URL bài báo")
            use_sample = gr.Checkbox(label="Sử dụng file mẫu (/mnt/data/baocaohocmay.docx)")

        with gr.Column():
            max_chunk_tokens = gr.Slider(128, 2048, value=800, step=64, label="Max chunk tokens")
            chunk_overlap = gr.Slider(0, 512, value=64, step=16, label="Chunk overlap (tokens)")
            max_summary_tokens = gr.Slider(32, 512, value=128, step=16, label="Max summary tokens per chunk")
            num_beams = gr.Slider(1, 8, value=4, step=1, label="Beams (num_beams)")

            btn = gr.Button("Tóm tắt ngay — Summarize")

    orig_out = gr.Textbox(label="Nội dung nguồn (cắt ngắn)", lines=10)
    chunk_summary_out = gr.Textbox(label="Kết quả tóm tắt (ghép từng chunk)", lines=15)
    short_out = gr.Textbox(label="Tóm tắt ngắn gọn", lines=10)

    btn.click(
        process_input,
        inputs=[model_dir, max_chunk_tokens, chunk_overlap, max_summary_tokens,
                num_beams, file, url, use_sample],
        outputs=[orig_out, chunk_summary_out, short_out]
    )

demo.launch(share=True)
